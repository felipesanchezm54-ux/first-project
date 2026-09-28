"""
Genera las tablas SEO y UX del taller a partir de una sola fuente (este archivo):

  docs/tabla-seo.md   docs/tabla-seo.csv
  docs/tabla-ux.md    docs/tabla-ux.csv
  docs/Tabla de parámetros SEO y UX - diligenciada.xlsx   (copia del Excel de la profesora con la columna llena)

Las filas siguen exactamente el orden del Excel original (01–15 SEO, 01–09 UX),
para poder copiar el CSV directo en la columna vacía.

Uso:  python3 scripts/generar_tablas.py      (requiere: pip install openpyxl)
"""

import csv
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
EXCEL_ORIGINAL = DOCS / "Tabla de parámetros SEO y UX.xlsx"
EXCEL_DILIGENCIADO = DOCS / "Tabla de parámetros SEO y UX - diligenciada.xlsx"

# ──────────────────────────────────────────────────────────────────────────
# Parámetros SEO — columna «✏️ CÓMO LO APLIQUÉ EN MI SITIO WEB»
# Cada fila: qué se hizo · dónde (página o archivo) · cómo se verifica.
# ──────────────────────────────────────────────────────────────────────────
SEO = [
    ("01", "E-E-A-T",
     "Qué: la experiencia y la autoría son visibles. /nosotros presenta al equipo con nombre, rol y especialidad (Felipe, Martín, Sofía, María Isabel y Vera, identificada como IA). "
     "Cada artículo del blog y el caso Tienda Óptica muestran autor, rol, fecha de publicación y fecha de actualización. "
     "El artículo de Habeas Data cita fuentes oficiales (Ley 1581 de 2012, Decreto 1377 de 2013, SIC). "
     "Dirección, teléfono y correo aparecen en el footer y en datos estructurados Organization/ProfessionalService. "
     "Confianza: testimonios y cifras simuladas llevan la etiqueta «ilustrativo» o «Datos de demostración». "
     "Dónde: app/nosotros/page.tsx, app/blog/[slug]/page.tsx, app/portafolio/tienda-optica/page.tsx, lib/jsonld.ts. "
     "Cómo se verifica: visitando las páginas y pasando la URL por el Rich Results Test de Google (schema Organization y Article)."),
    ("02", "Contenido útil y people-first",
     "Qué: cada página responde una sola intención de búsqueda, con una keyword por página (ver fila 03). "
     "Cada servicio explica qué incluye, qué se entrega, qué KPI se mide y un ejemplo real en Tienda Óptica. "
     "Los 4 artículos del blog (740–890 palabras) usan formato AEO: pregunta como H2 y respuesta directa de 40–59 palabras debajo, luego el desarrollo. "
     "No hay texto de relleno. "
     "Dónde: app/servicios/page.tsx, lib/content/services.ts, content/blog/*.mdx. "
     "Cómo se verifica: la prueba automática tests/seo-accesibilidad.spec.ts falla si aparece «lorem ipsum» o «Creative Nova»; también se puede leer cualquier página y comprobar que responde su pregunta en el primer párrafo."),
    ("03", "Keywords en ubicaciones prominentes",
     "Qué: la keyword principal está en el Title, el H1, la meta descripción, la URL cuando aplica, el primer párrafo y el alt de la imagen Open Graph. "
     "Inicio: «agencia de marketing digital en Medellín» (title, H1, meta). "
     "Nosotros: «equipo de marketing en Medellín». "
     "Servicios: «servicios de marketing digital para empresas». "
     "Planes: «planes de marketing digital… precios… Colombia». "
     "Blog: «blog de marketing digital». "
     "Portafolio: «casos de éxito en marketing digital». "
     "Caso: «estrategia de marketing para ópticas». "
     "Contacto: «contratar agencia de marketing en Medellín». "
     "Dónde: metadata y H1 de cada app/*/page.tsx. "
     "Cómo se verifica: ver el código fuente (Ctrl+U) y buscar la keyword en <title>, <meta name=\"description\"> y <h1>."),
    ("04", "Core Web Vitals",
     "Qué: el HTML se genera en el servidor (Next.js) y el CSS va incrustado en el HTML. "
     "Las fuentes se cargan con next/font y display: swap. "
     "No hay video en el hero: la animación es SVG y CSS. "
     "Los formularios (Zod y React Hook Form) se cargan en diferido; GA4 y Meta Pixel solo se cargan si el usuario acepta. "
     "Resultado en Lighthouse móvil (simulación 4G lenta): Performance 94–96. "
     "CLS: 0,035. TBT: 80–140 ms. LCP simulado: 2,5–2,9 s (real sin limitación de red: ~0,2 s en local). "
     "Dónde: next.config.ts (inlineCss), app/layout.tsx (next/font), components/lazy/. "
     "Cómo se verifica: docs/lighthouse/*.html o npm run lighthouse. Con el sitio publicado, en pagespeed.web.dev (meta: LCP < 2,5 s, INP < 200 ms, CLS < 0,1)."),
    ("05", "Mobile-First Index",
     "Qué: el diseño se hizo primero para móvil y se revisó en 375, 768 y 1440 px. "
     "El móvil muestra el mismo contenido que el escritorio, sin secciones ocultas. "
     "El texto base es de 16 px. Los CTA miden 48–56 px de alto y ningún botón baja de 44 px, el mínimo pedido. "
     "El menú hamburguesa es accesible y no hay scroll horizontal (se corrigió un desborde de tablas). "
     "Dónde: todas las páginas; components/layout/header.tsx (menú móvil). "
     "Cómo se verifica: capturas en docs/capturas/375, 768 y 1440; 30 pruebas del proyecto «movil» en Playwright; Lighthouse móvil."),
    ("06", "Title Tag",
     "Qué: un title único por página, con formato «Keyword principal | NEXO» y máximo 60 caracteres: "
     "Inicio «Agencia de marketing digital en Medellín | NEXO» (47) · "
     "Nosotros «Nosotros: equipo de marketing en Medellín | NEXO» (48) · "
     "Servicios «Servicios de marketing digital para empresas | NEXO» (51) · "
     "Planes «Planes de marketing digital y precios en Colombia | NEXO» (56) · "
     "Blog «Blog de marketing digital para marcas | NEXO» (44) · "
     "Portafolio «Casos de éxito en marketing digital | NEXO» (42) · "
     "Caso «Estrategia de marketing para ópticas | Caso NEXO» (48) · "
     "Contacto «Contratar agencia de marketing en Medellín | NEXO» (49). "
     "Cada artículo del blog tiene el suyo (45–50 caracteres). "
     "Dónde: export const metadata en cada app/*/page.tsx (helper lib/seo.ts). "
     "Cómo se verifica: pestaña del navegador o código fuente; la prueba automática exige 60 caracteres o menos y la marca en cada página."),
    ("07", "Meta Description",
     "Qué: una meta descripción única por página, de 146 a 153 caracteres (rango pedido: 120–155). "
     "Incluye la keyword y termina con un CTA claro: «Agenda tu asesoría gratuita», «Míralos», «Conócenos», «Compáralos y agenda tu asesoría», «Lee y aplica hoy». "
     "Ejemplo (Inicio, 146 caracteres): «Agencia de marketing digital en Medellín: estrategia, redes, contenido y pauta con un panel donde ves cada resultado. Agenda tu asesoría gratuita.» "
     "Dónde: campo description en el metadata de cada app/*/page.tsx y en el frontmatter de cada artículo MDX. "
     "Cómo se verifica: código fuente; la prueba automática exige 120–155 caracteres en las 15 páginas públicas."),
    ("08", "Estructura H1 / H2 / H3",
     "Qué: un solo H1 por página, con la keyword. Las secciones principales usan H2 (en el blog, cada H2 es una pregunta) y las tarjetas y subsecciones usan H3, sin saltar niveles. "
     "En el Inicio, el H1 combina la keyword y el mensaje: «Agencia de marketing digital en Medellín — Marketing que se mide». "
     "Dónde: todas las páginas; components/sections/section.tsx (SectionHeading usa H2). "
     "Cómo se verifica: la extensión HeadingsMap o la prueba automática, que exige un H1 por página y ningún salto de nivel (de H2 a H4, por ejemplo)."),
    ("09", "URL amigable",
     "Qué: URLs cortas, en minúscula, con guiones, sin tildes ni parámetros: "
     "/nosotros · /servicios · /planes · /blog · /portafolio · /portafolio/tienda-optica · /contacto. "
     "Los artículos llevan la keyword en la URL: /blog/embudo-de-marketing-digital-para-pymes, /blog/cuanto-invertir-en-meta-ads-al-empezar, /blog/habeas-data-para-pymes-ley-1581, /blog/cada-cuanto-publicar-en-redes-sociales. "
     "Los servicios usan anclas descriptivas (/servicios#publicidad-digital). "
     "Dónde: estructura de carpetas de app/ y nombres de archivo en content/blog/. "
     "Cómo se verifica: barra de direcciones y /sitemap.xml."),
    ("10", "Alt Text en imágenes",
     "Qué: el sitio usa ilustraciones SVG propias e íconos, no fotos de stock. "
     "Los elementos decorativos llevan aria-hidden y alt vacío, para que no sean ruido en lectores de pantalla. "
     "Las imágenes con significado tienen texto alternativo descriptivo: "
     "la imagen Open Graph («NEXO, agencia de marketing digital en Medellín», con keyword), "
     "los logos (aria-label «Logo de Tienda Óptica», «NEXO, ir al inicio») "
     "y las gráficas (role=img con descripción de los datos, por ejemplo «Barras mensuales de aumento de conversión: mes 1 +4,7 %…», más una tabla alternativa en el panel). "
     "Todo <img> futuro pasa por next/image, que exige alt. "
     "Dónde: app/opengraph-image.tsx, lib/seo.ts, components/ui/logo.tsx, app/portafolio/tienda-optica/page.tsx, components/panel/dashboard.tsx. "
     "Cómo se verifica: la prueba automática falla si existe un <img> sin alt; axe no reporta violaciones serias de «image-alt»."),
    ("11", "Links internos crawleables",
     "Qué: todos los enlaces internos usan <Link href=\"/ruta\"> (etiqueta <a> real, rastreable). "
     "Hay menú con 7 ítems y un desplegable a los 6 servicios, breadcrumbs en las páginas internas y footer con navegación y servicios. "
     "Además hay enlaces cruzados: servicio → caso Tienda Óptica, artículo → servicio relacionado, caso → blog de Habeas Data. "
     "Los textos de enlace son descriptivos («Leer el caso completo de Tienda Óptica», «Qué incluye redes sociales»); nunca «clic aquí». "
     "El sitemap tiene las 15 URLs públicas. "
     "Dónde: components/layout/header.tsx, footer.tsx, breadcrumbs.tsx, app/sitemap.ts. "
     "Cómo se verifica: la prueba automática exige al menos 2 enlaces internos en el contenido de cada página; el recorrido E2E hace clic en el menú, los CTA y las tarjetas, sin enlaces rotos (404 solo en la URL de prueba)."),
    ("12", "HTTPS / Seguridad SSL",
     "Qué: el sitio está listo para HTTPS. "
     "Envía los encabezados Strict-Transport-Security (HSTS), Content-Security-Policy, X-Content-Type-Options, X-Frame-Options, Referrer-Policy y Permissions-Policy. "
     "Canonical, sitemap y Open Graph usan URLs https://. "
     "La sesión del panel usa una cookie httpOnly, SameSite=Lax y Secure en producción, y las contraseñas se guardan cifradas con bcrypt. "
     "Al publicar en Vercel, el certificado SSL es automático. "
     "Dónde: next.config.ts (headers), lib/auth.ts. "
     "Cómo se verifica: candado en el navegador al publicar; curl -I https://dominio muestra los encabezados; Lighthouse Buenas prácticas: 100."),
    ("13", "Sin prácticas spam",
     "Qué: se usa una keyword por página, de forma natural, sin repetirla en exceso. "
     "Todo el contenido es original, escrito para NEXO. "
     "No hay texto oculto ni cloaking: el mismo HTML para Google y para las personas, y las FAQ que se cierran siguen en el HTML y coinciden con el schema FAQPage. "
     "No hay enlaces comprados. "
     "Las cifras simuladas están marcadas como tales. "
     "El formulario tiene campo trampa (honeypot) y límite de envíos para frenar el spam de bots. "
     "robots.txt excluye /panel y /api. "
     "Dónde: todo el sitio; app/robots.ts; app/api/contact/route.ts. "
     "Cómo se verifica: leer el contenido; en Search Console, sin acciones manuales; la prueba de FAQ confirma que la respuesta visible coincide con el schema."),
    ("14", "PageRank / Links entrantes",
     "Qué, dentro del sitio: el enlazado interno reparte autoridad hacia las páginas de negocio (servicios, caso, contacto). Las páginas pensadas para recibir enlaces (el caso Tienda Óptica y los 4 artículos) tienen una URL limpia e imagen Open Graph propia, para compartirse bien en redes. "
     "Plan de enlaces entrantes para cuando el sitio esté publicado (pendiente de ejecutar): "
     "(1) perfil de Google Business en Medellín; "
     "(2) perfiles de Instagram, Facebook, TikTok y LinkedIn enlazando al sitio (ya declarados en sameAs del JSON-LD); "
     "(3) crédito «Sitio y estrategia por NEXO» en la landing de Tienda Óptica; "
     "(4) directorios locales de empresas de Medellín; "
     "(5) mención del proyecto en canales de CEIPA. "
     "Dónde: lib/jsonld.ts (sameAs), lib/site.ts (redes). "
     "Cómo se verifica: en Search Console, informe «Enlaces», una vez el dominio esté verificado."),
    ("15", "Contenido fresco y actualizado",
     "Qué: los artículos y el caso muestran la fecha de actualización real (por ejemplo, «Actualizado el 24 de septiembre de 2026»), además de dateModified en el JSON-LD Article y lastModified en el sitemap. "
     "Las cifras del caso se recalcularon con los datos del panel (+17,9 %, 1.130 leads). "
     "El blog tiene 4 categorías para publicar de forma periódica (el plan Full Brand incluye 2 artículos al mes). "
     "/llms.txt se genera solo con la lista actualizada de artículos. "
     "Regla del equipo: solo se cambia la fecha cuando cambia el contenido. "
     "Dónde: frontmatter updatedAt en content/blog/*.mdx, lib/content/cases.ts, app/sitemap.ts, app/llms.txt/route.ts. "
     "Cómo se verifica: ver las fechas en las páginas, /sitemap.xml y el schema Article en el Rich Results Test."),
]

# ──────────────────────────────────────────────────────────────────────────
# Leyes de UX — columna «✏️ CÓMO LO APLIQUÉ EN MI SITIO»
# ──────────────────────────────────────────────────────────────────────────
UX = [
    ("01", "Ley de Fitts",
     "Qué: los CTA tienen un área táctil de 48 px de alto (56 px en los CTA grandes) y ocupan todo el ancho en móvil; los botones secundarios compactos (header, cookies, filtros del panel) miden 44 px, el mínimo de la tabla. "
     "El CTA principal («Potencia tu marca») está en el hero, visible sin scroll, y tiene efecto magnético con mouse. "
     "Cada sección termina con su CTA donde acaba la lectura. "
     "WhatsApp y Vera son botones flotantes abajo a la derecha, en la zona del pulgar. No se superponen entre sí, suben cuando aparece el aviso de cookies y se ocultan en móvil mientras escribes en un formulario, para no tapar el botón «Enviar». "
     "Dónde: components/ui/button.tsx (min-h-12 / min-h-14), components/layout/floating-actions.tsx, app/page.tsx. "
     "Cómo se verifica: capturas en docs/capturas/375; prueba manual en celular; en el código, los CTA usan min-h-12 o min-h-14 (48 o 56 px) y ningún botón baja de min-h-11 (44 px)."),
    ("02", "Ley de Hick",
     "Qué: el menú principal tiene 7 ítems (Inicio, Nosotros, Servicios, Planes, Blog, Portafolio, Contacto). "
     "Hay 3 planes, con Growth resaltado como recomendado. "
     "Cada sección tiene un solo CTA primario (verde); los secundarios van con contorno o como enlace. "
     "El formulario de contacto está dividido en 2 pasos, y el del inicio pide solo correo y servicio. "
     "Dónde: lib/site.ts (mainNav), components/sections/pricing.tsx, components/forms/contact-form.tsx. "
     "Cómo se verifica: contar los ítems del menú (7) y los CTA primarios por sección (1)."),
    ("03", "Ley de Semejanza",
     "Qué: hay un solo estilo de botón primario en todo el sitio (verde teal, texto oscuro, forma de píldora), definido en un componente. "
     "Las 6 tarjetas de servicio tienen la misma estructura (ícono → título → frase → enlace). "
     "Los enlaces dentro de un texto van siempre subrayados y del mismo color. "
     "Los íconos son de una sola familia (Lucide), con los mismos tamaños. "
     "Dónde: components/ui/button.tsx, components/sections/service-card.tsx, app/globals.css (.link, .prose-nexo a). "
     "Cómo se verifica: revisión visual en docs/capturas; todos los botones de acción se ven iguales."),
    ("04", "Ley de Miller",
     "Qué: la información va en bloques de 4 a 7 elementos: 6 servicios, 5 valores, 5 bandas de lead scoring, 6 KPI, 4 etapas del embudo y 3 columnas por servicio (qué incluye, entregables, KPI). "
     "Los artículos usan párrafos cortos bajo H2 y H3, con tabla de contenido. "
     "El formulario agrupa los campos por tema: «Tus datos» y «Tu proyecto». "
     "Una barra muestra el paso en que vas. "
     "Dónde: app/servicios/page.tsx, components/panel/dashboard.tsx, components/forms/contact-form.tsx. "
     "Cómo se verifica: contar los elementos de cada bloque y escanear una página por sus títulos."),
    ("05", "Ley de Proximidad",
     "Qué: cada etiqueta va pegada a su campo, con la ayuda y el error justo debajo (componente Field). "
     "El precio va junto al nombre del plan y lo que incluye. "
     "Cada cifra va pegada a su explicación («+17,9 %, conversión digital, vs. línea base, mes 5 de 6»). "
     "El CTA «Solicitar [servicio]» aparece inmediatamente después de la descripción de cada servicio. "
     "Dónde: components/ui/field.tsx, components/sections/metric-grid.tsx, components/sections/pricing.tsx, app/servicios/page.tsx. "
     "Cómo se verifica: revisión visual; los elementos relacionados están juntos."),
    ("06", "Ley de Tesler",
     "Qué: la complejidad la asume el sistema. "
     "WhatsApp abre con el mensaje ya escrito. "
     "Elegir un plan lleva a /contacto?plan=… con el plan preseleccionado, y «Solicitar [servicio]» preselecciona el servicio. "
     "Los campos tienen autocomplete del navegador (nombre, correo, teléfono, empresa). "
     "El texto de ejemplo del mensaje cambia según el servicio elegido. "
     "En el login, el botón «Usar la cuenta demo» llena los datos. "
     "Vera interpreta las métricas del panel por el cliente. "
     "El lenguaje es sin jerga (por ejemplo, «CPL = costo por lead»). "
     "Dónde: lib/site.ts (whatsappUrl), components/forms/contact-form.tsx, components/panel/login-form.tsx, lib/dashboard.ts (veraInsight). "
     "Cómo se verifica: ir de Planes → «Agenda una asesoría» y ver el plan ya elegido en el paso 2 (hay una prueba automática)."),
    ("07", "Ley de Postel",
     "Qué: el sistema es flexible con lo que recibe. "
     "Acepta el teléfono con o sin +57, con espacios o guiones, y el correo en mayúsculas; los guarda normalizados (+57XXXXXXXXXX y minúsculas). "
     "Los errores dicen qué falló y cómo corregirlo, junto al campo: «Correo inválido: verifica que incluya @ y un dominio, por ejemplo nombre@empresa.com», «Teléfono inválido: escribe 10 dígitos, con o sin +57». "
     "La 404 explica qué pudo pasar y ofrece 4 salidas útiles. "
     "En el formulario se puede volver al paso anterior sin perder datos. "
     "Dónde: lib/schemas.ts (normalizePhone, mensajes), app/not-found.tsx. "
     "Cómo se verifica: probar el formulario con datos incorrectos; la prueba automática revisa los mensajes de nombre, correo y autorización, y que el correo en mayúsculas y el teléfono con +57 y espacios se acepten."),
    ("08", "Ley de Zeigarnik",
     "Qué: el formulario de contacto muestra el progreso («Paso 1 de 2», con barra) y termina con una confirmación clara: «¡Mensaje enviado! Te contactamos en menos de 24 horas hábiles.» "
     "El blog tiene barra de progreso de lectura bajo el header. "
     "El panel muestra el avance hacia la meta (+17,9 % de +20 %, barra al 89 %). "
     "El flujo de email de carrito abandonado aplica la misma ley con los clientes de Tienda Óptica. "
     "Dónde: components/forms/contact-form.tsx, components/blog/reading-progress.tsx, components/panel/dashboard.tsx. "
     "Cómo se verifica: llenar el formulario hasta el final y ver la confirmación; hacer scroll en un artículo."),
    ("09", "Ley de Jakob",
     "Qué: el sitio se comporta como otros sitios conocidos. "
     "El logo está arriba a la izquierda y lleva al inicio. "
     "El acceso «Panel de clientes» está arriba a la derecha (donde se espera el login o carrito). "
     "Hay menú hamburguesa en móvil, footer con enlaces legales y enlaces de texto subrayados. "
     "Los acordeones, filtros y el chat funcionan como en cualquier sitio. "
     "Dónde: components/layout/header.tsx, footer.tsx. "
     "Cómo se verifica: navegar el sitio; la prueba E2E recorre el menú en escritorio y la hamburguesa en móvil."),
]

HEAD_SEO = ["#", "Parámetro", "Cómo lo apliqué en mi sitio web"]
HEAD_UX = ["#", "Ley de UX", "Cómo lo apliqué en mi sitio"]


def write_md(path: Path, title: str, intro: str, head: list[str], rows):
    lines = [f"# {title}", "", intro, "", f"| {' | '.join(head)} |", f"|{'---|' * len(head)}"]
    for n, name, how in rows:
        lines.append(f"| {n} | **{name}** | {how.replace('|', '/')} |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_csv(path: Path, head: list[str], rows):
    # UTF-8 con BOM para que Excel muestre bien las tildes; comillas en todos los campos.
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL)
        w.writerow(head)
        w.writerows(rows)


def fill_excel():
    try:
        import openpyxl
        from openpyxl.styles import Alignment
    except ImportError:
        print("openpyxl no está instalado: se omite la copia diligenciada del Excel (pip install openpyxl).")
        return
    shutil.copyfile(EXCEL_ORIGINAL, EXCEL_DILIGENCIADO)
    wb = openpyxl.load_workbook(EXCEL_DILIGENCIADO)
    for sheet, rows in (("Parámetros SEO", SEO), ("Parámetros UX", UX)):
        ws = wb[sheet]
        by_num = {n: how for n, _, how in rows}
        for r in range(1, ws.max_row + 1):
            num = ws.cell(r, 1).value
            if isinstance(num, str) and num in by_num:
                assert ws.cell(r, 2).value.strip().lower().startswith(dict((n, name) for n, name, _ in rows)[num].lower()[:6]), (sheet, num)
                c = ws.cell(r, 7, by_num[num])
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["G"].width = 90
    wb.save(EXCEL_DILIGENCIADO)
    print(f"✓ {EXCEL_DILIGENCIADO.relative_to(ROOT)}")


if __name__ == "__main__":
    write_md(DOCS / "tabla-seo.md", "Tabla de parámetros SEO — cómo los apliqué en el sitio de NEXO",
             "Mismo orden que la hoja «Parámetros SEO» del Excel del taller. Cada fila dice **qué** se hizo, **dónde** (página o archivo) y **cómo se verifica**. "
             "Generada con `python3 scripts/generar_tablas.py`.", HEAD_SEO, SEO)
    write_md(DOCS / "tabla-ux.md", "Tabla de leyes de UX — cómo las apliqué en el sitio de NEXO",
             "Mismo orden que la hoja «Parámetros UX» del Excel del taller. Cada fila dice **qué** se hizo, **dónde** y **cómo se verifica**. "
             "Generada con `python3 scripts/generar_tablas.py`.", HEAD_UX, UX)
    write_csv(DOCS / "tabla-seo.csv", HEAD_SEO, SEO)
    write_csv(DOCS / "tabla-ux.csv", HEAD_UX, UX)
    print("✓ docs/tabla-seo.md, docs/tabla-seo.csv, docs/tabla-ux.md, docs/tabla-ux.csv")
    fill_excel()
